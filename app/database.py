
import sqlite3, hashlib, shutil, json, traceback, os, re
from pathlib import Path
from datetime import datetime, timedelta
from contextlib import contextmanager

ROOT=Path(__file__).resolve().parent.parent
DATA_DIR=ROOT/"data"; BACKUP_DIR=ROOT/"backups"; LOG_DIR=ROOT/"logs"
DB_PATH=DATA_DIR/"ranisaa_erp.db"; KEY_FILE=DATA_DIR/".backup_key"
for p in [DATA_DIR,BACKUP_DIR,LOG_DIR]: p.mkdir(exist_ok=True)
SCHEMA_VERSION=5

def log_error(msg):
    try:(LOG_DIR/"error.log").open("a",encoding="utf-8").write(f"\n[{datetime.now().isoformat()}] {msg}\n")
    except:pass

def get_db_connection():
    c=sqlite3.connect(DB_PATH,timeout=30); c.row_factory=sqlite3.Row
    c.execute("PRAGMA journal_mode=WAL"); c.execute("PRAGMA synchronous=FULL"); c.execute("PRAGMA foreign_keys=ON")
    return c

@contextmanager
def tx():
    c=get_db_connection()
    try: yield c; c.commit()
    except Exception: c.rollback(); raise
    finally:c.close()

def _ensure_label_columns():
    with tx() as c:
        pcols=[r[1] for r in c.execute("PRAGMA table_info(parties)").fetchall()]
        if "manufacturer_code" not in pcols:
            c.execute("ALTER TABLE parties ADD COLUMN manufacturer_code TEXT DEFAULT ''")
        icols=[r[1] for r in c.execute("PRAGMA table_info(items)").fetchall()]
        if "color" not in icols:
            c.execute("ALTER TABLE items ADD COLUMN color TEXT DEFAULT ''")


def h(p):return hashlib.sha256(p.encode()).hexdigest()

def init_database():
    with tx() as c:
        c.execute("CREATE TABLE IF NOT EXISTS schema_meta(key TEXT PRIMARY KEY,value TEXT)")
        c.execute("""CREATE TABLE IF NOT EXISTS users(id INTEGER PRIMARY KEY AUTOINCREMENT,username TEXT UNIQUE NOT NULL,
            password_hash TEXT NOT NULL,role TEXT NOT NULL,full_name TEXT,created_at TEXT DEFAULT CURRENT_TIMESTAMP,last_login TEXT)""")
        c.execute("""CREATE TABLE IF NOT EXISTS role_permissions(role TEXT,module TEXT,allowed INTEGER,PRIMARY KEY(role,module))""")
        c.execute("""CREATE TABLE IF NOT EXISTS parties(id INTEGER PRIMARY KEY AUTOINCREMENT,name TEXT UNIQUE NOT NULL,gstin TEXT DEFAULT '',contact TEXT DEFAULT '',address TEXT DEFAULT '')""")
        c.execute("""CREATE TABLE IF NOT EXISTS customers(id INTEGER PRIMARY KEY AUTOINCREMENT,name TEXT NOT NULL,mobile TEXT UNIQUE NOT NULL,address TEXT DEFAULT '',created_at TEXT DEFAULT CURRENT_TIMESTAMP,updated_at TEXT DEFAULT CURRENT_TIMESTAMP)""")
        c.execute("""CREATE TABLE IF NOT EXISTS hsn_codes(code TEXT PRIMARY KEY,description TEXT DEFAULT '')""")
        c.execute("""CREATE TABLE IF NOT EXISTS items(id INTEGER PRIMARY KEY AUTOINCREMENT,barcode TEXT UNIQUE NOT NULL,name TEXT NOT NULL,hsn_code TEXT DEFAULT '',
            design_number TEXT DEFAULT '',purchase_price REAL DEFAULT 0,mrp REAL DEFAULT 0,sale_price REAL DEFAULT 0,gst_rate REAL DEFAULT 0,
            unit TEXT DEFAULT 'PCS',stock_qty REAL DEFAULT 0,min_stock REAL DEFAULT 0,category TEXT DEFAULT '',brand TEXT DEFAULT '',
            created_at TEXT DEFAULT CURRENT_TIMESTAMP,updated_at TEXT DEFAULT CURRENT_TIMESTAMP)""")
        for _col, _typ in (("color","TEXT DEFAULT ''"),("size","TEXT DEFAULT ''"),("attribute_type","TEXT DEFAULT ''"),("attribute_detail","TEXT DEFAULT ''")):
            try: c.execute(f"ALTER TABLE items ADD COLUMN {_col} {_typ}")
            except Exception: pass
        c.execute("""CREATE TABLE IF NOT EXISTS purchases(id INTEGER PRIMARY KEY AUTOINCREMENT,bill_number TEXT UNIQUE NOT NULL,
            supplier_name TEXT,supplier_gstin TEXT DEFAULT '',bill_date DATE,total_amount REAL DEFAULT 0,discount REAL DEFAULT 0,
            tax_amount REAL DEFAULT 0,net_amount REAL DEFAULT 0,payment_mode TEXT DEFAULT 'CASH',notes TEXT,created_by INTEGER,created_at TEXT DEFAULT CURRENT_TIMESTAMP)""")
        c.execute("""CREATE TABLE IF NOT EXISTS purchase_items(id INTEGER PRIMARY KEY AUTOINCREMENT,purchase_id INTEGER,item_id INTEGER,
            quantity REAL,rate REAL,discount REAL DEFAULT 0,gst_rate REAL,amount REAL,hsn_code TEXT,design_number TEXT)""")
        for _col, _typ in (("color","TEXT DEFAULT ''"),("size","TEXT DEFAULT ''"),("attribute_type","TEXT DEFAULT ''"),("attribute_detail","TEXT DEFAULT ''")):
            try: c.execute(f"ALTER TABLE purchase_items ADD COLUMN {_col} {_typ}")
            except Exception: pass
        c.execute("""CREATE TABLE IF NOT EXISTS old_stock_entries(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            design_number TEXT NOT NULL,
            party_code TEXT NOT NULL,
            quantity REAL NOT NULL DEFAULT 1,
            rate REAL NOT NULL DEFAULT 0,
            gst_rate REAL NOT NULL DEFAULT 0,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP
        )""")
        for _col, _typ in (("color","TEXT DEFAULT ''"),("size","TEXT DEFAULT ''"),("attribute_type","TEXT DEFAULT ''"),("attribute_detail","TEXT DEFAULT ''")):
            try: c.execute(f"ALTER TABLE old_stock_entries ADD COLUMN {_col} {_typ}")
            except Exception: pass
        c.execute("""CREATE TABLE IF NOT EXISTS sales(id INTEGER PRIMARY KEY AUTOINCREMENT,bill_number TEXT UNIQUE NOT NULL,
            customer_name TEXT,customer_mobile TEXT,customer_address TEXT,bill_date DATE,bill_time TIME,subtotal REAL DEFAULT 0,
            discount REAL DEFAULT 0,tax_amount REAL DEFAULT 0,total_amount REAL DEFAULT 0,payment_mode TEXT DEFAULT 'CASH',
            status TEXT DEFAULT 'COMPLETED',created_by INTEGER,created_at TEXT DEFAULT CURRENT_TIMESTAMP)""")
        c.execute("""CREATE TABLE IF NOT EXISTS sale_items(id INTEGER PRIMARY KEY AUTOINCREMENT,sale_id INTEGER,item_id INTEGER,
            quantity REAL,rate REAL,discount REAL DEFAULT 0,gst_rate REAL,amount REAL)""")
        c.execute("""CREATE TABLE IF NOT EXISTS returns(id INTEGER PRIMARY KEY AUTOINCREMENT,return_number TEXT UNIQUE NOT NULL,
            original_bill_number TEXT,customer_name TEXT,customer_mobile TEXT,return_date DATE,return_time TIME,reason TEXT,
            total_amount REAL DEFAULT 0,created_by INTEGER,created_at TEXT DEFAULT CURRENT_TIMESTAMP)""")
        c.execute("""CREATE TABLE IF NOT EXISTS return_items(id INTEGER PRIMARY KEY AUTOINCREMENT,return_id INTEGER,item_id INTEGER,
            quantity REAL,rate REAL,amount REAL)""")
        c.execute("""CREATE TABLE IF NOT EXISTS settings(key TEXT PRIMARY KEY,value TEXT DEFAULT '')""")
        c.execute("""CREATE TABLE IF NOT EXISTS backup_log(id INTEGER PRIMARY KEY AUTOINCREMENT,backup_file TEXT,backup_type TEXT,file_size INTEGER,status TEXT,created_at TEXT DEFAULT CURRENT_TIMESTAMP)""")
        c.execute("""CREATE TABLE IF NOT EXISTS audit_log(id INTEGER PRIMARY KEY AUTOINCREMENT,user_id INTEGER,action TEXT,entity TEXT,entity_id TEXT,details TEXT,created_at TEXT DEFAULT CURRENT_TIMESTAMP)""")
        defaults={"company_name":"RANISAA EXCLUSIVE","bill_prefix":"BILL","purchase_prefix":"PUR","return_prefix":"RET","bill_template_path":"","barcode_prefix":"RAN","barcode_start":"100001",
                  "auto_backup":"1","backup_retention_days":"30","app_version":"1.0.0","data_version":str(SCHEMA_VERSION)}
        for k,v in defaults.items(): c.execute("INSERT OR IGNORE INTO settings VALUES(?,?)",(k,v))
        if not c.execute("SELECT 1 FROM users LIMIT 1").fetchone():
            c.execute("INSERT INTO users(username,password_hash,role,full_name) VALUES(?,?,?,?)",("admin",h("admin123"),"admin","Administrator"))
            c.execute("INSERT INTO users(username,password_hash,role,full_name) VALUES(?,?,?,?)",("staff",h("staff123"),"staff","Staff"))
        mods=["purchase","stock","mrp","billing","reports","returns","customers","parties","settings","backup","users","billview"]
        for m in mods:c.execute("INSERT OR IGNORE INTO role_permissions VALUES('admin',?,1)",(m,))
        st={"purchase":1,"stock":1,"mrp":0,"billing":1,"reports":1,"returns":1,"customers":1,"parties":1,"settings":0,"backup":0,"users":0,"billview":0}
        for m,v in st.items():c.execute("INSERT OR IGNORE INTO role_permissions VALUES('staff',?,?)",(m,v))
        c.execute("INSERT OR REPLACE INTO schema_meta VALUES('version',?)",(str(SCHEMA_VERSION),))

def setting(k,d=""):
    with get_db_connection() as c:
        r=c.execute("SELECT value FROM settings WHERE key=?",(k,)).fetchone();return r["value"] if r else d
def set_setting(k,v):
    with tx() as c:c.execute("INSERT INTO settings VALUES(?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",(k,str(v)))
def login(u,p):
    with get_db_connection() as c:
        r=c.execute("SELECT * FROM users WHERE username=? AND password_hash=?",(u,h(p))).fetchone()
        return dict(r) if r else None
def perms(role):
    with get_db_connection() as c:return {r["module"]:bool(r["allowed"]) for r in c.execute("SELECT module,allowed FROM role_permissions WHERE role=?",(role,))}
def set_perm(role,module,v):
    with tx() as c:c.execute("INSERT INTO role_permissions VALUES(?,?,?) ON CONFLICT(role,module) DO UPDATE SET allowed=excluded.allowed",(role,module,int(v)))

def party(name,gstin="",manufacturer_code=""):
    with tx() as c:
        cols=[r[1] for r in c.execute("PRAGMA table_info(parties)").fetchall()]
        if "manufacturer_code" in cols:
            c.execute("""INSERT INTO parties(name,gstin,manufacturer_code) VALUES(?,?,?)
                         ON CONFLICT(name) DO UPDATE SET gstin=excluded.gstin,manufacturer_code=excluded.manufacturer_code""",
                      (name.strip(),gstin.strip().upper(),manufacturer_code.strip().upper()))
        else:
            c.execute("""INSERT INTO parties(name,gstin) VALUES(?,?)
                         ON CONFLICT(name) DO UPDATE SET gstin=excluded.gstin""",
                      (name.strip(),gstin.strip().upper()))
def parties():
    with get_db_connection() as c:return [dict(r) for r in c.execute("SELECT * FROM parties ORDER BY name")]
def party_by_name(n):
    with get_db_connection() as c:
        r=c.execute("SELECT * FROM parties WHERE lower(name)=lower(?)",(n.strip(),)).fetchone();return dict(r) if r else None

def save_customer(name,mobile,address=""):
    if not mobile:return
    with tx() as c:c.execute("""INSERT INTO customers(name,mobile,address) VALUES(?,?,?)
        ON CONFLICT(mobile) DO UPDATE SET name=excluded.name,address=excluded.address,updated_at=CURRENT_TIMESTAMP""",(name.strip(),mobile.strip(),address))
def customer_by_mobile(mobile):
    with get_db_connection() as c:
        r=c.execute("SELECT * FROM customers WHERE mobile=?",(mobile.strip(),)).fetchone();return dict(r) if r else None
def customers():
    with get_db_connection() as c:return [dict(r) for r in c.execute("SELECT * FROM customers ORDER BY name")]

def hsn(q=""):
    with get_db_connection() as c:return [r["code"] for r in c.execute("SELECT code FROM hsn_codes WHERE code LIKE ? ORDER BY code",(q.strip()+"%",))]
def save_hsn(code):
    with tx() as c:c.execute("INSERT OR IGNORE INTO hsn_codes(code) VALUES(?)",(code.strip(),))

def save_hsn_code(code):
    save_hsn(code)

def purchase_exists(bill_number, supplier_name):
    with get_db_connection() as c:
        return c.execute("SELECT 1 FROM purchases WHERE bill_number=? AND lower(supplier_name)=lower(?) LIMIT 1",(bill_number,supplier_name)).fetchone() is not None

def purchase_lookup(party_name="", bill_number=""):
    with get_db_connection() as c:
        q="SELECT id,bill_number,supplier_name,supplier_gstin,bill_date,total_amount FROM purchases WHERE 1=1"; args=[]
        if party_name: q+=" AND lower(supplier_name)=lower(?)"; args.append(party_name)
        if bill_number: q+=" AND bill_number=?"; args.append(bill_number)
        return [dict(r) for r in c.execute(q+" ORDER BY bill_date DESC,id DESC",args)]

def purchase_items_lookup(purchase_id):
    with get_db_connection() as c:
        return [dict(r) for r in c.execute("""SELECT pi.item_id,pi.quantity,pi.rate,pi.gst_rate,pi.hsn_code,pi.design_number,
            pi.color,pi.size,pi.attribute_type,pi.attribute_detail,
            i.barcode,i.name,i.mrp,i.sale_price,i.color,i.size,i.attribute_type FROM purchase_items pi JOIN items i ON i.id=pi.item_id
            WHERE pi.purchase_id=? ORDER BY pi.id""",(purchase_id,))]

def barcode_design_search(design="", party="", invoice="", bill_date=""):
    q=str(design or "").strip()
    party=str(party or "").strip()
    invoice=str(invoice or "").strip()
    bill_date=str(bill_date or "").strip()
    with get_db_connection() as c:
        sql="""SELECT i.id AS item_id,i.barcode,i.name,i.hsn_code,i.design_number,
                      i.mrp,i.purchase_price,i.gst_rate,i.stock_qty,i.color,i.size,i.attribute_type,
                      p.bill_number,p.bill_date,p.supplier_name,
                      pa.manufacturer_code
               FROM items i
               JOIN purchase_items pi ON pi.item_id=i.id
               JOIN purchases p ON p.id=pi.purchase_id
               LEFT JOIN parties pa ON lower(pa.name)=lower(p.supplier_name)
               WHERE 1=1"""
        args=[]
        if q:
            sql+=" AND i.design_number LIKE ?"; args.append("%"+q+"%")
        if party:
            sql+=" AND lower(p.supplier_name)=lower(?)"; args.append(party)
        if invoice:
            sql+=" AND p.bill_number=?"; args.append(invoice)
        if bill_date:
            sql+=" AND p.bill_date=?"; args.append(bill_date)
        sql+=" ORDER BY i.design_number,p.id DESC"
        return [dict(r) for r in c.execute(sql,args).fetchall()]

def barcode_old_stock_search(design="", party=""):
    q=str(design or "").strip()
    party=str(party or "").strip()
    with get_db_connection() as c:
        sql="""SELECT os.id AS old_id,os.design_number,os.party_code,os.quantity,
                      os.rate,os.gst_rate,os.color,os.size,os.attribute_type,
                      i.id AS item_id,i.barcode,i.name,i.mrp,i.stock_qty
               FROM old_stock_entries os
               LEFT JOIN items i ON i.design_number=os.design_number
                                  AND i.brand=os.party_code
                                  AND i.color=os.color
                                  AND i.size=os.size
               WHERE 1=1"""
        args=[]
        if q:
            sql+=" AND os.design_number LIKE ?"; args.append("%"+q+"%")
        if party:
            sql+=" AND upper(os.party_code)=upper(?)"; args.append(party)
        sql+=" ORDER BY os.design_number,os.id DESC"
        return [dict(r) for r in c.execute(sql,args).fetchall()]

def next_barcode():
    prefix=setting("barcode_prefix","RAN"); n=int(setting("barcode_start","100001"))
    with tx() as c:
        while c.execute("SELECT 1 FROM items WHERE barcode=?",(f"{prefix}{n}",)).fetchone(): n+=1
        set_setting("barcode_start",str(n+1))
        return f"{prefix}{n}"

def set_item_barcode(item_id,barcode):
    with tx() as c:c.execute("UPDATE items SET barcode=?,updated_at=CURRENT_TIMESTAMP WHERE id=?",(str(barcode),int(item_id)))

def set_item_mrp(item_id,mrp):
    with tx() as c:c.execute("UPDATE items SET mrp=?,updated_at=CURRENT_TIMESTAMP WHERE id=?",(float(mrp),item_id))


def _parse_attribute_values(attribute_type, attribute_values):
    """Parse clothing attributes without changing the existing stock model."""
    attr_type=str(attribute_type or "").strip().upper()
    raw=str(attribute_values or "").strip()
    if attr_type not in ("","NONE","COLOR","SIZE","BOTH"):
        raise ValueError("Attribute must be COLOR, SIZE or BOTH.")

    def clean(values):
        return list(dict.fromkeys(v.strip().upper() for v in values if v.strip()))

    if attr_type in ("","NONE"):
        return raw.upper(), "", []
    if attr_type=="COLOR":
        colors=clean(re.split(r"[,/&]+",raw))
        return ", ".join(colors), "", colors
    if attr_type=="SIZE":
        sizes=clean(re.split(r"[,/&]+",raw))
        return "", ", ".join(sizes), sizes

    # BOTH format:
    # BLUE L,XL / GREEN XXL
    colors=[]
    mappings=[]
    details=[]
    for group in [g.strip() for g in raw.split("/") if g.strip()]:
        bits=group.split(None,1)
        if len(bits)!=2:
            raise ValueError("BOTH format: BLUE L,XL / GREEN XXL")
        color=bits[0].strip().upper()
        group_sizes=clean(re.split(r"[,]+",bits[1]))
        if not group_sizes:
            raise ValueError("BOTH format: BLUE L,XL / GREEN XXL")
        colors.append(color)
        mappings.append(f"{color}: {', '.join(group_sizes)}")
        details.append({"color":color,"sizes":group_sizes})
    colors=clean(colors)
    return ", ".join(colors), " / ".join(mappings), details


def add_old_stock(design_number, party_code, quantity, rate, gst_rate,
                  attribute_type="", attribute_values=""):
    design_number=str(design_number).strip().upper()
    party_code=str(party_code).strip().upper()
    quantity=float(quantity); rate=float(rate); gst_rate=float(gst_rate)
    if not design_number or not party_code:
        raise ValueError("Design Number and Party Code are required.")
    if quantity <= 0:
        raise ValueError("Quantity must be greater than zero.")
    if gst_rate not in (5.0,12.0,18.0):
        raise ValueError("Old Stock GST must be 5%, 12% or 18%.")

    color,size,detail=_parse_attribute_values(attribute_type,attribute_values)
    attribute_detail=json.dumps(detail, ensure_ascii=False) if isinstance(detail,list) else str(detail or "")
    attr_type=str(attribute_type or "").strip().upper()
    if attr_type=="NONE": attr_type=""

    with tx() as c:
        existing=c.execute(
            """SELECT id,quantity FROM old_stock_entries
               WHERE design_number=? AND party_code=? AND rate=? AND gst_rate=? AND color=? AND size=?""",
            (design_number,party_code,rate,gst_rate,color,size)
        ).fetchone()

        if existing:
            c.execute("UPDATE old_stock_entries SET quantity=quantity+? WHERE id=?",
                      (quantity,existing["id"]))
            oid=int(existing["id"])
        else:
            c.execute(
                """INSERT INTO old_stock_entries
                   (design_number,party_code,quantity,rate,gst_rate,color,size,attribute_type,attribute_detail)
                   VALUES(?,?,?,?,?,?,?,?,?)""",
                (design_number,party_code,quantity,rate,gst_rate,color,size,attr_type,attribute_detail)
            )
            oid=int(c.execute("SELECT last_insert_rowid()").fetchone()[0])

        item=c.execute(
            """SELECT id FROM items
               WHERE design_number=? AND brand=? AND color=? AND size=?
               ORDER BY id LIMIT 1""",
            (design_number,party_code,color,size)
        ).fetchone()

        if item:
            c.execute(
                """UPDATE items SET stock_qty=stock_qty+?,purchase_price=?,sale_price=?,
                   gst_rate=?,attribute_type=?,attribute_detail=?,updated_at=CURRENT_TIMESTAMP WHERE id=?""",
                (quantity,rate,rate,gst_rate,attr_type,attribute_detail,item["id"])
            )
        else:
            barcode=f"OLD-{oid:06d}"
            while c.execute("SELECT 1 FROM items WHERE barcode=?",(barcode,)).fetchone():
                oid+=1; barcode=f"OLD-{oid:06d}"
            c.execute(
                """INSERT INTO items
                   (barcode,name,hsn_code,design_number,purchase_price,mrp,sale_price,gst_rate,stock_qty,
                    category,brand,color,size,attribute_type,attribute_detail)
                   VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (barcode,"OLD STOCK","",design_number,rate,0,rate,gst_rate,quantity,
                 "OLD STOCK",party_code,color,size,attr_type,attribute_detail)
            )
    return oid

def old_stock_entries():
    with get_db_connection() as c:
        rows=c.execute("""SELECT * FROM old_stock_entries ORDER BY id DESC""").fetchall()
        return [dict(r) for r in rows]

def old_stock_by_item(item_id):
    with get_db_connection() as c:
        r=c.execute("""SELECT os.* FROM old_stock_entries os
                       JOIN items i ON i.design_number=os.design_number
                                      AND i.brand=os.party_code
                       WHERE i.id=? ORDER BY os.id DESC LIMIT 1""",(item_id,)).fetchone()
        return dict(r) if r else None

def items():
    with get_db_connection() as c:return [dict(r) for r in c.execute("SELECT * FROM items ORDER BY name")]
def item_barcode(b):
    with get_db_connection() as c:
        r=c.execute("SELECT * FROM items WHERE barcode=?",(b,)).fetchone();return dict(r) if r else None
def item_source_purchase(item_id):
    with get_db_connection() as c:
        r=c.execute("""SELECT
            i.id AS item_id,i.barcode,i.name,i.hsn_code,i.design_number,i.mrp,i.sale_price,
            i.gst_rate,i.stock_qty,
            p.id AS purchase_id,p.bill_number,p.bill_date,p.supplier_name,p.supplier_gstin,p.total_amount
            FROM items i
            LEFT JOIN purchase_items pi ON pi.item_id=i.id
            LEFT JOIN purchases p ON p.id=pi.purchase_id
            WHERE i.id=?
            ORDER BY p.bill_date DESC,p.id DESC
            LIMIT 1""",(int(item_id),)).fetchone()
        return dict(r) if r else None

def item_search(q):
    with get_db_connection() as c:return [dict(r) for r in c.execute("SELECT * FROM items WHERE name LIKE ? OR barcode LIKE ? OR hsn_code LIKE ? OR design_number LIKE ? ORDER BY name",(f"%{q}%",)*4)]
def add_item(name,barcode,hsn_code,design,purchase,gst):
    with tx() as c:
        c.execute("""INSERT INTO items(barcode,name,hsn_code,design_number,purchase_price,gst_rate) VALUES(?,?,?,?,?,?)
                     ON CONFLICT(barcode) DO UPDATE SET name=excluded.name,hsn_code=excluded.hsn_code,design_number=excluded.design_number,
                     purchase_price=excluded.purchase_price,gst_rate=excluded.gst_rate,updated_at=CURRENT_TIMESTAMP""",(barcode,name,hsn_code,design,purchase,gst))
        return c.execute("SELECT id FROM items WHERE barcode=?",(barcode,)).fetchone()["id"]

def next_no(kind):
    key={"sale":"bill_prefix","purchase":"purchase_prefix","return":"return_prefix"}[kind];prefix=setting(key);k=key+"_n";n=int(setting(k,"1"))
    table={"sale":"sales","purchase":"purchases","return":"returns"}[kind];col="bill_number" if kind!="return" else "return_number"
    with tx() as c:
        while c.execute(f"SELECT 1 FROM {table} WHERE {col}=?",(f"{prefix}{n:06d}",)).fetchone():n+=1
        set_setting(k,n+1);return f"{prefix}{n:06d}"

def delete_party(name):
    with tx() as c:
        r=c.execute("SELECT id FROM parties WHERE lower(name)=lower(?)",(name.strip(),)).fetchone()
        if not r:
            raise ValueError("Party not found")
        # Do not delete historical purchase records. Remove only the party master entry.
        c.execute("DELETE FROM parties WHERE id=?",(r["id"],))

def save_purchase_atomic(data, rows):
    if not rows:
        raise ValueError("No purchase items.")
    bill=str(data["bill"]).strip()
    party_name=str(data["party"]).strip()
    if not bill:
        raise ValueError("Invoice number is required.")
    if not party_name:
        raise ValueError("Party name is required.")

    with tx() as c:
        if c.execute("SELECT 1 FROM purchases WHERE bill_number=? LIMIT 1",(bill,)).fetchone():
            raise ValueError(f"Purchase invoice {bill} already exists.")

        # Every item, purchase header, purchase rows and stock update are in ONE transaction.
        # If anything fails, tx() rolls everything back.
        staged=[]
        for i,x in enumerate(rows,1):
            name=str(x.get("name","")).strip()
            hsn=str(x.get("hsn","")).strip()
            design=str(x.get("design","")).strip()
            color=str(x.get("color","")).strip().upper()
            size=str(x.get("size","")).strip().upper()
            attr_type=str(x.get("attribute_type","") or "").strip().upper()
            if attr_type=="COLOR":
                color=", ".join(dict.fromkeys(v.strip().upper() for v in re.split(r"[,/&]+",color) if v.strip())); size=""
            elif attr_type=="SIZE":
                size=", ".join(dict.fromkeys(v.strip().upper() for v in re.split(r"[,/&]+",size) if v.strip())); color=""
            elif attr_type=="BOTH":
                color=", ".join(dict.fromkeys(v.strip().upper() for v in str(x.get("color","")).split(",") if v.strip()))
                size=", ".join(dict.fromkeys(v.strip().upper() for v in str(x.get("size","")).split(",") if v.strip()))
            else:
                color=""; size=""; attr_type=""
            attr_detail=" / ".join(x.get("attribute_detail") or [])
            qty=float(x.get("qty",0))
            price=float(x.get("price",0))
            gst=float(x.get("gst",0))
            if not name or not hsn or not design or qty <= 0 or price <= 0:
                raise ValueError(f"Invalid item data at row {i}.")
            # Unique barcode for the staged item; it can later be replaced by the MRP/barcode module.
            import uuid
            barcode="AUTO-"+uuid.uuid4().hex.upper()
            item_cursor = c.execute(
                """INSERT INTO items(barcode,name,hsn_code,design_number,purchase_price,mrp,sale_price,gst_rate,stock_qty,color,size,attribute_type,attribute_detail)
                   VALUES(?,?,?,?,?,?,?,?,0,?,?,?,?)""",
                (barcode,name,hsn,design,price,0,0,gst,color,size,attr_type,str(x.get("attribute_detail", "")))
            )
            iid=item_cursor.lastrowid
            staged.append((iid,x))

        c.execute(
            """INSERT INTO parties(name,gstin) VALUES(?,?)
               ON CONFLICT(name) DO UPDATE SET gstin=excluded.gstin""",
            (party_name,str(data.get("gstin","")).strip().upper())
        )

        purchase_cursor = c.execute(
            """INSERT INTO purchases
               (bill_number,supplier_name,supplier_gstin,bill_date,total_amount,tax_amount,net_amount,payment_mode,created_by)
               VALUES(?,?,?,?,?,?,?,?,?)""",
            (bill,party_name,str(data.get("gstin","")).strip().upper(),data["date"],
             float(data["total"]),float(data["gst"]),float(data["total"]),
             "CASH",data["user"])
        )
        pid=purchase_cursor.lastrowid

        for iid,x in staged:
            taxable=float(x["tax"])
            gst_amount=float(x["ga"])
            qty=float(x["qty"])
            price=float(x["price"])
            gst=float(x["gst"])
            c.execute(
                """INSERT INTO purchase_items
                   (purchase_id,item_id,quantity,rate,discount,gst_rate,amount,hsn_code,design_number,color,size,attribute_type)
                   VALUES(?,?,?,?,?,?,?,?,?,?,?,?)""",
                (pid,iid,qty,price,0,gst,taxable+gst_amount,str(x["hsn"]),str(x["design"]),
                 str(x.get("color","")),
                 str(x.get("size","")),str(x.get("attribute_type","")))
            )
            c.execute(
                """UPDATE items
                   SET stock_qty=stock_qty+?,purchase_price=?,gst_rate=?,hsn_code=?,design_number=?,
                       color=?,size=?,attribute_type=?,attribute_detail=?,updated_at=CURRENT_TIMESTAMP
                   WHERE id=?""",
                (qty,price,gst,str(x["hsn"]),str(x["design"]),
                 str(x.get("color","")),str(x.get("size","")),str(x.get("attribute_type","")),str(x.get("attribute_detail","")),iid)
            )

        return pid

def purchase(data,rows):
    with tx() as c:
        c.execute("""INSERT INTO purchases(bill_number,supplier_name,supplier_gstin,bill_date,total_amount,tax_amount,net_amount,created_by)
                     VALUES(?,?,?,?,?,?,?,?)""",(data["bill"],data["party"],data["gstin"],data["date"],data["total"],data["gst"],data["total"],data["user"]))
        pid=c.lastrowid
        for x in rows:
            c.execute("INSERT INTO purchase_items(purchase_id,item_id,quantity,rate,gst_rate,amount,hsn_code,design_number) VALUES(?,?,?,?,?,?,?,?)",(pid,x["id"],x["qty"],x["price"],x["gst"],x["tax"],x["hsn"],x["design"]))
            c.execute("UPDATE items SET stock_qty=stock_qty+?,purchase_price=?,gst_rate=?,hsn_code=?,design_number=?,updated_at=CURRENT_TIMESTAMP WHERE id=?",(x["qty"],x["price"],x["gst"],x["hsn"],x["design"],x["id"]))
        return pid

def sale(data,rows):
    with tx() as c:
        for x in rows:
            r=c.execute("SELECT stock_qty FROM items WHERE id=?",(x["id"],)).fetchone()
            if not r or r["stock_qty"]<x["qty"]:raise ValueError("Insufficient stock")
        c.execute("""INSERT INTO sales(bill_number,customer_name,customer_mobile,customer_address,bill_date,bill_time,subtotal,discount,tax_amount,total_amount,payment_mode,created_by)
                     VALUES(?,?,?,?,?,?,?,?,?,?,?,?)""",(data["bill"],data["name"],data["mobile"],data["address"],data["date"],data["time"],data["subtotal"],data["discount"],data["gst"],data["total"],data["pay"],data["user"]))
        sid=c.lastrowid
        for x in rows:
            c.execute("INSERT INTO sale_items(sale_id,item_id,quantity,rate,discount,gst_rate,amount) VALUES(?,?,?,?,?,?,?)",(sid,x["id"],x["qty"],x["rate"],x["disc"],x["gst"],x["amount"]))
            c.execute("UPDATE items SET stock_qty=stock_qty-?,updated_at=CURRENT_TIMESTAMP WHERE id=?",(x["qty"],x["id"]))
        save_customer(data["name"],data["mobile"],data["address"])
        return sid

def sales_report(a,b):
    with get_db_connection() as c:return [dict(r) for r in c.execute("""SELECT bill_number,bill_date,bill_time,customer_name,customer_mobile,subtotal,discount,tax_amount,total_amount,payment_mode
        FROM sales WHERE bill_date BETWEEN ? AND ? ORDER BY bill_date,bill_time""",(a,b))]
def purchases_report(a,b):
    with get_db_connection() as c:return [dict(r) for r in c.execute("SELECT bill_number,bill_date,supplier_name,tax_amount,total_amount FROM purchases WHERE bill_date BETWEEN ? AND ? ORDER BY bill_date",(a,b))]


def users_list():
    with get_db_connection() as c:return [dict(r) for r in c.execute("SELECT id,username,role,full_name,created_at,last_login FROM users ORDER BY id")]
def update_user_credentials(user_id,username,password=None,full_name=None,role=None):
    with tx() as c:
        old=c.execute("SELECT * FROM users WHERE id=?",(user_id,)).fetchone()
        if not old: raise ValueError("User not found")
        if password:
            c.execute("UPDATE users SET username=?,password_hash=?,full_name=COALESCE(?,full_name),role=COALESCE(?,role) WHERE id=?",
                      (username.strip(),h(password),full_name,role,user_id))
        else:
            c.execute("UPDATE users SET username=?,full_name=COALESCE(?,full_name),role=COALESCE(?,role) WHERE id=?",
                      (username.strip(),full_name,role,user_id))

def sale_by_bill(bill):
    with get_db_connection() as c:
        r=c.execute("SELECT * FROM sales WHERE bill_number=?",(bill,)).fetchone()
        return dict(r) if r else None
def purchase_by_bill(bill):
    with get_db_connection() as c:
        r=c.execute("SELECT * FROM purchases WHERE bill_number=?",(bill,)).fetchone()
        return dict(r) if r else None
def sale_items(sale_id):
    with get_db_connection() as c:
        return [dict(r) for r in c.execute("""SELECT si.*,i.name,i.barcode,i.hsn_code,i.mrp
            FROM sale_items si JOIN items i ON i.id=si.item_id WHERE si.sale_id=? ORDER BY si.id""",(sale_id,))]
def purchase_items(purchase_id):
    return purchase_items_lookup(purchase_id)
def recent_bills(kind):
    table="sales" if kind=="SALE" else "purchases"
    with get_db_connection() as c:
        return [dict(r) for r in c.execute(f"SELECT id,bill_number,{'customer_name' if kind=='SALE' else 'supplier_name'} AS party,bill_date,total_amount FROM {table} ORDER BY id DESC LIMIT 500")]

def _excel_snapshot(path):
    """Export every current business table to one XLSX workbook.
    This is a reporting/archive copy; the SQLite database remains the authoritative source.
    """
    from openpyxl import Workbook
    path=Path(path); path.parent.mkdir(parents=True, exist_ok=True)
    with get_db_connection() as c:
        tables=[r[0] for r in c.execute("SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%' ORDER BY name")]
        wb=Workbook(); wb.remove(wb.active)
        for name in tables:
            ws=wb.create_sheet((name[:31] or 'DATA'))
            rows=c.execute(f'SELECT * FROM "{name.replace(chr(34), chr(34)*2)}"').fetchall()
            cols=[d[0] for d in c.execute(f'SELECT * FROM "{name.replace(chr(34), chr(34)*2)}" LIMIT 0').description]
            if cols: ws.append(cols)
            for row in rows: ws.append([row[col] for col in cols])
            ws.freeze_panes='A2'
            ws.auto_filter.ref=ws.dimensions if ws.max_row and ws.max_column else 'A1'
            for col in ws.columns:
                width=min(max(len(str(cell.value or '')) for cell in col)+2, 45)
                ws.column_dimensions[col[0].column_letter].width=width
        wb.save(path)
    return path

def _encrypt_file(src, enc):
    from cryptography.fernet import Fernet
    if not KEY_FILE.exists(): KEY_FILE.write_bytes(Fernet.generate_key())
    enc=Path(enc); enc.write_bytes(Fernet(KEY_FILE.read_bytes()).encrypt(Path(src).read_bytes()))
    return enc

def encrypted_backup(include_excel=True):
    """Create an encrypted SQLite snapshot and, optionally, an encrypted full-data Excel archive."""
    try:
        from cryptography.fernet import Fernet
        if not KEY_FILE.exists(): KEY_FILE.write_bytes(Fernet.generate_key())
        stamp=datetime.now().strftime('%Y%m%d_%H%M%S')
        raw=BACKUP_DIR/f'db_{stamp}.sqlite'; enc=BACKUP_DIR/f'db_{stamp}.sqlite.enc'
        src=get_db_connection(); dst=sqlite3.connect(raw)
        try: src.backup(dst)
        finally: dst.close(); src.close()
        _encrypt_file(raw,enc); raw.unlink(missing_ok=True)
        with tx() as c:
            c.execute("INSERT INTO backup_log(backup_file,backup_type,file_size,status) VALUES(?,?,?,?)",(enc.name,'AUTO_ENCRYPTED_DB',enc.stat().st_size,'SUCCESS'))
        excel_enc=None
        if include_excel:
            xraw=BACKUP_DIR/f'full_data_{stamp}.xlsx'; xenc=BACKUP_DIR/f'full_data_{stamp}.xlsx.enc'
            _excel_snapshot(xraw); _encrypt_file(xraw,xenc); xraw.unlink(missing_ok=True)
            with tx() as c:
                c.execute("INSERT INTO backup_log(backup_file,backup_type,file_size,status) VALUES(?,?,?,?)",(xenc.name,'AUTO_ENCRYPTED_EXCEL',xenc.stat().st_size,'SUCCESS'))
            excel_enc=xenc
        # Keep the existing configured retention policy for both backup types.
        try:
            days=max(1,int(setting('backup_retention_days','30')))
            cutoff=datetime.now()-timedelta(days=days)
            for f in BACKUP_DIR.glob('*.enc'):
                try:
                    if datetime.fromtimestamp(f.stat().st_mtime) < cutoff: f.unlink()
                except Exception: pass
        except Exception: pass
        return {'db':enc,'excel':excel_enc}
    except Exception as e:
        log_error('backup '+traceback.format_exc()); return None

def restore(path):
    try:
        from cryptography.fernet import Fernet
        p=Path(path);raw=Fernet(KEY_FILE.read_bytes()).decrypt(p.read_bytes());tmp=DB_PATH.with_suffix('.tmp');tmp.write_bytes(raw)
        c=sqlite3.connect(tmp);ok=c.execute('PRAGMA integrity_check').fetchone()[0]=='ok';c.close()
        if not ok: tmp.unlink(missing_ok=True); return False
        if DB_PATH.exists(): shutil.copy2(DB_PATH,DB_PATH.with_suffix('.pre_restore.bak'))
        shutil.move(tmp,DB_PATH);return True
    except Exception as e:log_error('restore '+traceback.format_exc());return False

def database_integrity():
    try:
        with get_db_connection() as c: return c.execute('PRAGMA integrity_check').fetchone()[0]=='ok'
    except Exception as e: log_error('integrity '+traceback.format_exc()); return False

def auto_repair():
    """Safe repair: run additive schema setup, verify integrity, and make an encrypted backup.
    It never rewrites/deletes business rows and does not pretend to repair arbitrary Python bugs.
    """
    init_database()
    if not database_integrity(): raise RuntimeError('DATABASE INTEGRITY CHECK FAILED')
    return encrypted_backup(include_excel=True)

def auto_backup():
    if setting('auto_backup','1')!='1': return None
    with get_db_connection() as c:
        r=c.execute("SELECT created_at FROM backup_log WHERE status='SUCCESS' AND backup_type='AUTO_ENCRYPTED_DB' ORDER BY id DESC LIMIT 1").fetchone()
    if not r or datetime.fromisoformat(r['created_at']).date()!=datetime.now().date():
        return encrypted_backup(include_excel=True)
    return None

try:
    _ensure_label_columns()
except Exception:
    pass
